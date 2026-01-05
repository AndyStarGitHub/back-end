from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from app.schemas.quiz import (
    QuizCreate,
    QuizQuestionCreate,
    QuizAnswerOptionCreate
)

from app.utils.quiz_import import (
    canon_header,
    norm_text,
    parse_bool,
    parse_uuid,
    parse_frequency,
)

from app.schemas.quiz_import import (
    ImportErrorItem,
    ParsedQuizItem,
    ParsedImportResult,
)


EXPECTED_COLS = {
    "quizid",
    "quiztitle",
    "quizdescription",
    "quizfrequency",
    "questiontitle",
    "optiontext",
    "iscorrect",
}


def parse_quizzes_from_excel(
        file_bytes: bytes,
        *,
        max_rows: int,
) -> ParsedImportResult:
    f = io.BytesIO(file_bytes)
    wb = load_workbook(f, data_only=True)

    sheet = wb["Quizzes"] if "Quizzes" in wb.sheetnames else wb.worksheets[0]

    if sheet.max_row and sheet.max_row > max_rows:
        return ParsedImportResult(
            items=[],
            errors=[
                ImportErrorItem(
                    row=1,
                    field="file",
                    message=f"Excel has too many rows: {sheet.max_row} > {max_rows}",
                )
            ],
            total_quizzes_in_file=0,
        )

    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return ParsedImportResult(
            items=[],
            errors=[ImportErrorItem(
                row=1,
                field="file",
                message="Empty sheet"
            )],
            total_quizzes_in_file=0,
        )

    col_map: dict[str, int] = {}
    for idx, col in enumerate(header):
        if col is None:
            continue
        name = canon_header(str(col))
        col_map[name] = idx

    missing = [c for c in EXPECTED_COLS if c not in col_map]
    if missing:
        return ParsedImportResult(
            items=[],
            errors=[
                ImportErrorItem(
                    row=1,
                    field="header",
                    message=f"Missing columns: {', '.join(missing)}",
                )
            ],
            total_quizzes_in_file=0,
        )

    quizzes: dict[str, dict] = {}
    errors: list[ImportErrorItem] = []

    def get_cell(row_vals: tuple[Any, ...], name: str) -> Any:
        return row_vals[col_map[name]] \
            if col_map[name] < len(row_vals) \
            else None

    excel_row_num = 1
    error_message = "Required (must be filled on every row; merged/blank cells are not supported)"
    for row_vals in rows_iter:
        excel_row_num += 1
        if row_vals is None:
            continue

        try:
            quiz_id = parse_uuid(get_cell(row_vals, "quizid"))
        except Exception as e:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "quiz_id",
                    error_message,
                )
            )
            continue

        quiz_title = str(get_cell(row_vals, "quiztitle") or "").strip()
        if not quiz_title:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "quiz_title",
                    error_message,
                )
            )
            continue

        quiz_desc_val = get_cell(row_vals, "quizdescription")
        quiz_description = None \
            if quiz_desc_val is None \
            else str(quiz_desc_val).strip() or None

        try:
            frequency = parse_frequency(get_cell(row_vals, "quizfrequency"))
        except Exception as e:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "quiz_frequency",
                    error_message,
                )
            )
            continue

        question_title = str(get_cell(row_vals, "questiontitle") or "").strip()
        if not question_title:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "question_title",
                    error_message,
                )
            )
            continue

        option_text = str(get_cell(row_vals, "optiontext") or "").strip()
        if not option_text:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "option_text",
                    error_message,
                )
            )
            continue

        try:
            is_correct = parse_bool(get_cell(row_vals, "iscorrect"))
        except Exception as e:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "is_correct",
                    error_message,
                )
            )

            continue

        quiz_title_norm = norm_text(quiz_title)
        question_title_norm = norm_text(question_title)
        option_text_norm = norm_text(option_text)

        if quiz_id is not None:
            quiz_key = f"id:{str(quiz_id)}"
        else:
            quiz_key = f"title:{quiz_title_norm}"

        if quiz_key not in quizzes:
            quizzes[quiz_key] = {
                "quiz_id": quiz_id,
                "quiz_title": quiz_title,
                "quiz_title_norm": quiz_title_norm,
                "quiz_description": quiz_description,
                "frequency": frequency,
                "first_row": excel_row_num,
                "questions": {},
            }

        qmap: dict[str, dict] = quizzes[quiz_key]["questions"]
        if question_title_norm not in qmap:
            qmap[question_title_norm] = {
                "title": question_title,
                "options": {},
            }

        opt_map: dict[str, tuple[str, bool, int]] = qmap[question_title_norm]["options"]
        if option_text_norm in opt_map:
            errors.append(
                ImportErrorItem(
                    excel_row_num,
                    "option_text",
                    f"Duplicate option_text in the same question: {option_text!r}",
                )
            )
            continue

        opt_map[option_text_norm] = (option_text, is_correct, excel_row_num)

    items: list[ParsedQuizItem] = []
    for qk, data in quizzes.items():
        questions: list[QuizQuestionCreate] = []
        for _, qdata in data["questions"].items():
            options: list[QuizAnswerOptionCreate] = [
                QuizAnswerOptionCreate(text=txt, is_correct=isc)
                for (txt, isc, _rownum) in qdata["options"].values()
            ]
            questions.append(
                QuizQuestionCreate(
                    title=qdata["title"],
                    options=options,
                )
            )

        quiz_create = QuizCreate(
            title=data["quiz_title"],
            description=data["quiz_description"],
            frequency=data["frequency"],
            questions=questions,
        )

        items.append(
            ParsedQuizItem(
                quiz_id=data["quiz_id"],
                title=data["quiz_title"],
                description=data["quiz_description"],
                frequency=data["frequency"],
                quiz_title_norm=data["quiz_title_norm"],
                quiz_create=quiz_create,
                first_row=data["first_row"],
            )
        )

    return ParsedImportResult(
        items=items,
        errors=errors,
        total_quizzes_in_file=len(items),
    )
